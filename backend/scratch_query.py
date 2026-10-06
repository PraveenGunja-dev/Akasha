import sys; sys.path.append('d:/Akasha/backend')
from pptx import Presentation
from pptx.oxml.ns import qn
import shutil

# Start with a fresh template
template = r"D:\Akasha\backend\Data\uploads\03. Mundra North_CPAG_Energy_Monthly Report_Aug'26.pptx"
output = r"D:\Akasha\backend\cache\cpag_wind\test.pptx"
shutil.copyfile(template, output)

prs = Presentation(output)
chart = None
for sh in prs.slides[6].shapes:
    if getattr(sh, "has_chart", False):
        chart = sh.chart
        break

# Dummy data
cats = ["Apr-25", "May-25", "Jun-25", "Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26", "Apr-26", "May-26", "Jun-26", "Jul-26"]
data = [0.01, 0.02, 0.03, 0.04, 0.05, 0.06, 0.07, 0.08, 0.09, 0.10, 0.11, 0.12, 0.13, 0.14, 0.15, 0.16]

series_keys = ["monthly_plan", "monthly_actual", "monthly_forecast", "cum_plan", "cum_actual", "cum_forecast"]

for i, s in enumerate(chart.series):
    # 1. Update Categories
    cat = s._element.xpath('.//c:cat')
    if cat:
        caches = cat[0].xpath('.//c:strCache | .//c:numCache | .//c:strLit | .//c:numLit')
        if caches:
            pts = caches[0].findall(qn("c:pt"))
            for j, cat_val in enumerate(cats):
                if j < len(pts):
                    v = pts[j].find(qn("c:v"))
                    if v is not None:
                        v.text = str(cat_val)

    # 2. Update Values
    val = s._element.xpath('.//c:val')
    if val:
        caches = val[0].xpath('.//c:numCache | .//c:numLit')
        if caches:
            pts = caches[0].findall(qn("c:pt"))
            for j, pt_val in enumerate(data):
                if j < len(pts):
                    v = pts[j].find(qn("c:v"))
                    if v is not None:
                        v.text = str(pt_val)

prs.save(output)
print("Saved to", output)
