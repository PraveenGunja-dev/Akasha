"""
The CPAG pack as page images - the on-screen pack is the downloadable file.

The deck is built once per distinct input (data + manual entries + builder),
saved, and rendered page by page through PowerPoint. The viewer shows those
pages and its Download serves that same saved file, so what is reviewed on
screen and what is sent to management cannot differ.
"""
import hashlib
import json
import subprocess
import threading
from pathlib import Path
from typing import Any, Callable, Dict, List

from pptx import Presentation

BACKEND = Path(__file__).resolve().parent.parent
CACHE = BACKEND / "cache" / "cpag"
SCRIPT = BACKEND / "scripts" / "render_pptx.ps1"
_BUILDER_FILES = [BACKEND / "services" / "cpag_pptx.py", BACKEND / "services" / "cpag_render.py",
                  BACKEND / "assets" / "cpag_reference.pptx"]
_lock = threading.Lock()


def deck_key(data: Dict[str, Any]) -> str:
    h = hashlib.sha1()
    h.update(json.dumps(data, sort_keys=True, default=str).encode())
    for f in _BUILDER_FILES:
        h.update(str(f.stat().st_mtime_ns).encode())
    return h.hexdigest()[:16]


_STAR = "* after a date = P6 forecast (the milestone has not happened yet)."

# Where every value on a page comes from, keyed by the page's title. Shown in
# the viewer's info panel so a reviewer never has to ask.
SOURCES = [
    ("physical progress", "phase1", [
        "No P6 schedule exists for this Phase-1 unit, so there is nothing to compute.",
        "Table and graph are manual entry.",
    ]),
    ("physical progress", "scurve", [
        "P6 weightage units (the Nonlabor resource each schedule is loaded with = 100%).",
        "Plan: the P6 re-baseline B2 (B1 where the project has no B2), each activity's units spread over its planned duration - the method that reproduces the approved pack's plan line.",
        "Actual: P6 actual units up to the P6 data date, spread over each activity's actual dates.",
        "FTM = the month of the P6 data date. The table's Total row is the graph's value for that month.",
        "Wtg. = each area's share of the schedule. Variance = Plan - Actual.",
        "Remarks are manual entry - blank until the reviewer writes one.",
    ]),
    ("engineering progress", "", [
        "The project team's Master Document List (MDL Excel), uploaded here - no copy is kept beyond the most recent upload, so re-uploading a newer MDL is how this page refreshes.",
        "P6 holds only the engineering deliverable activities (29 per project, e.g. PSS-11), not a document register, so it cannot fill this page.",
        "J (Cum. Total approved) and K (Approved %) are calculated from the uploaded Cat I-IV* counts. Remarks come from the MDL's own Remarks column.",
    ]),
    ("procurement status", "", ["Section divider from the approved template."]),
    ("procurement", "", [
        "Packages through PO Date: the WBS/ZPS021 mapping uploaded on this page (BESS PMAG, 2026-09-22/26) - Package, Manufacturer, UOM, Scope and PO Date all come from that upload, never from P6 or the live SAP sync. One PO under a package is one row; CSS has no WBS mapping yet and stays on P6, marked manual provision until it does.",
        "Expected Delivery Start/Finish, MDCC Date, Delivered At Site and the forecast months stay P6 (Procurement > Ordering & Delivery): Start/Finish = each 'Receipt at Site' lot, MDCC Date = latest 'MDCC LOT', forecast months = lots whose forecast receipt falls in that month.",
        "CDD (Commercial Delivery Date) = SAP PO delivery date; no connected extract carries that field, so it shows '-'.",
        "Remarks are manual entry - blank until the reviewer writes one.",
        _STAR,
    ]),
    ("civil construction", "", [
        "P6 Material resources (installed quantity) under Construction, by element and stage, as of the P6 data date.",
        "Plan = quantity due by the data date on the plan baseline; Actual = P6 actual quantity. Stages P6 counts in sub-units (piles, m3) are shown in the element's own unit.",
        "Progress % = P6 weightage of the element's activities. Colours: 80%+ of plan green, 40-80% yellow, below 40% red.",
    ]),
    ("construction progress", "", [
        "P6 Material resources (installed quantity) under Construction, by element and stage, as of the P6 data date.",
        "Plan = quantity due by the data date on the plan baseline; Actual = P6 actual quantity.",
        "Progress % = P6 weightage of the element's activities. Colours: 80%+ of plan green, 40-80% yellow, below 40% red.",
        "'-' = this project's P6 has no such work.",
    ]),
    ("electrical", "", [
        "P6 Material resources under Construction: Scope, Plan (quantity due by the P6 data date on the plan baseline) and Actual.",
        "Plan % and Actual % are of scope. '-' = this project's P6 has no such resource.",
    ]),
    ("mandays", "", [
        "P6 Labor units per project, stacked.",
        "Plan: today's labour units phased on the plan baseline's activity dates.",
        "Actual: planned units x % complete, spread over actual dates - P6 does not post actual labour units; it reduces remaining units as work progresses.",
    ]),
    ("manpower deployment", "", [
        "Mandays divided by the days in the month - the same ratio the approved pack uses.",
        "Source as the Mandays page: P6 Labor units.",
    ]),
    ("contractor manpower", "", [
        "Manual entry - weekly forecast and actual headcount per contractor. No connected system holds it.",
        "Shortfall, totals and past-month average are calculated from the entries.",
    ]),
    ("critical issues", "", ["Manual entry - the issue register has no system of record."]),
    ("financial s curve", "", ["Manual entry from the EAC sheet until the automated EAC link is live."]),
    ("ordering status", "", [
        "P6 (Procurement WBS): package, quantity, Release of PO / SO ('Placement of the order') and the 'Order placed' update.",
        "Specification, TBER, vendors, offer and NFA dates are manual entry. 'NA' in Verified by Engineering consultant is the pack's convention.",
        _STAR,
    ]),
    ("approvals", "", [
        "P6 Statutory WBS: activity names, baseline and actual/forecast dates. Where the projects of a group differ, the project finishing last is shown.",
        "Responsible and Approval Authority: from the approved pack where the activity matches, otherwise manual entry.",
        "Remarks are manual entry - blank until the reviewer writes one.",
        _STAR,
    ]),
    ("commissioning plan", "", [
        "Manual entry - no connected system holds the month-wise MWh and container commissioning plan.",
        "Project and battery OEM headings are from the project register.",
    ]),
    ("salient features", "", [
        "Project Capacity = sum of the declared capacities.",
        "NFA table, Capex / PPA / Tariff text, Connectivity, Land Lease and Contractors rows are the approved pack's declared facts (editable by manual entry). SPV from the project register.",
    ]),
    ("overall layout", "", ["Kept exactly as the approved template."]),
    ("project capacity", "", ["Declared capacities from the approved CPAG pack - no system of record holds them."]),
    ("battery energy storage systems", "", ["Report date = the day the pack is generated. Layout from the approved template."]),
]


def _source(title: str, section_page: bool) -> List[str]:
    t = title.lower()
    if section_page:
        return ["Section divider from the approved template."]
    if t.startswith("thank you") or "growth" in t:
        return ["Closing page from the approved template."]
    for key, kind, lines in SOURCES:
        if key in t:
            if kind == "phase1" and not any(k in t for k in ("10a", "5a", "8a")):
                continue
            if kind == "scurve" and any(k in t for k in ("10a", "5a", "8a")):
                continue
            return lines
    return ["Layout from the approved template."]


def _outline(pptx_path: Path) -> List[Dict[str, Any]]:
    """Title and section for each page, for the viewer's jump list. A
    section-divider page starts a new section, as it does in the pack."""
    prs = Presentation(str(pptx_path))
    out, section = [], "Overview"
    for i, s in enumerate(prs.slides, start=1):
        heads = sorted((sh for sh in s.shapes
                        if sh.has_text_frame and sh.text_frame.text.strip()
                        and not sh.name.startswith("Slide Number")
                        and (sh.top or 0) < 914400 * 1.2),
                       key=lambda sh: sh.top or 0)
        if "Section Header" in s.slide_layout.name:
            heads = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()
                     and not sh.name.startswith("Slide Number")]
        title = (heads[0].text_frame.text.strip().split("\n")[0] if heads
                 else ("Cover" if i == 1 else "Project Capacity" if i == 2 else f"Page {i}"))
        title = title.replace("\n", " ")
        is_section = "Section Header" in s.slide_layout.name
        if is_section:
            section = title
        out.append({"n": i, "title": title[:120], "section": section,
                    "source": _source(title, is_section)})
    return out


def render(data: Dict[str, Any], build: Callable[[Dict[str, Any]], bytes]) -> Dict[str, Any]:
    key = deck_key(data)
    folder = CACHE / key
    meta_file = folder / "outline.json"
    with _lock:
        if not meta_file.exists():
            folder.mkdir(parents=True, exist_ok=True)
            deck = folder / "deck.pptx"
            deck.write_bytes(build(data))
            subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                 "-File", str(SCRIPT), "-Deck", str(deck), "-OutDir", str(folder)],
                check=True, capture_output=True, timeout=900)
            meta_file.write_text(json.dumps(_outline(deck)))
    return {"key": key, "slides": json.loads(meta_file.read_text())}


def page_path(key: str, n: int) -> Path:
    return CACHE / key / f"s{n:03d}.png"


def deck_path(key: str) -> Path:
    return CACHE / key / "deck.pptx"
