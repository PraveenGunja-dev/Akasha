"""
CPAG deck composition — the pack's full slide sequence.

The circulated pack has a fixed running order, and reviewers read it by slide
position. So the sequence is reproduced whole. Where no connected system holds
a slide's data, the slide is still issued — titled, numbered and in place —
carrying a statement of what is missing and what would fill it. Dropping those
slides would make the deck look like the pack had shrunk, when the reporting
gap is itself the finding.

Selecting one project shortens the deck (the per-project sections collapse from
six to one) but every slide that does appear is built the same way and carries
the same columns.
"""
from typing import Any, Dict, List, Optional

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from services.cpag_pptx_base import (
    BAND, _generated_by, BODY_W, CONTENT_TOP, INK, MARGIN, MUTED, RULE,
    _content_slide, _date, _heading, _footnote, _layout, _line_chart,
    _bar_chart, _month, _num, _pct, _section_slide, _table, _table_grouped,
    _text, _variance_tint, _drop_empty_placeholders,
    BRAND_BLUE, BRAND_PURPLE, WHITE, SLIDE_W, SLIDE_H, COVER_PHOTO,
    GREY_COL, PEACH_COL,
)


def _next_months(last_month: Optional[str], n: int) -> List[str]:
    """The next `n` month strings ("YYYY-MM") after `last_month`, for the
    pack's "Forecast Delivery Schedule" columns. Falls back to a fixed pair
    when no reported month is known, rather than leaving the header blank."""
    if not last_month or "-" not in last_month:
        return ["2026-09", "2026-10"][:n]
    y, m = (int(x) for x in last_month.split("-")[:2])
    out = []
    for _ in range(n):
        m += 1
        if m > 12:
            m = 1
            y += 1
        out.append(f"{y:04d}-{m:02d}")
    return out


def _last_reported_month(project: Optional[Dict[str, Any]]) -> Optional[str]:
    curve = (project or {}).get("sCurve") or []
    reported = [c["month"] for c in curve if c.get("actualCumPct") is not None]
    return max(reported) if reported else None


# ── Slides that exist but cannot be populated ──────────────────────────────
def _unavailable(prs, title: str, subtitle: str, reason: str,
                 needs: str = "") -> None:
    slide = _content_slide(prs, title, subtitle)
    top = CONTENT_TOP + Inches(1.5)
    panel = slide.shapes.add_shape(
        1, MARGIN + Inches(1.4), top, BODY_W - Inches(2.8), Inches(2.0))
    panel.fill.solid()
    panel.fill.fore_color.rgb = BAND
    panel.line.color.rgb = RULE
    panel.shadow.inherit = False

    _text(slide, MARGIN + Inches(1.8), top + Inches(0.34),
          BODY_W - Inches(3.6), Inches(0.3),
          "NOT AVAILABLE FROM CONNECTED SYSTEMS",
          size=11, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    _text(slide, MARGIN + Inches(1.8), top + Inches(0.80),
          BODY_W - Inches(3.6), Inches(0.7), reason,
          size=12, color=INK, align=PP_ALIGN.CENTER)
    if needs:
        _text(slide, MARGIN + Inches(1.8), top + Inches(1.48),
              BODY_W - Inches(3.6), Inches(0.42), needs,
              size=10, color=MUTED, italic=True, align=PP_ALIGN.CENTER)


def _cover(prs, title: str, subtitle: str, stamp: str = "") -> None:
    slide = prs.slides.add_slide(_layout(prs, "Blank"))
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xD9, 0xEE, 0xF7)

    if COVER_PHOTO.exists():
        slide.shapes.add_picture(
            str(COVER_PHOTO),
            left=Inches(0),
            top=Inches(0.9),
            width=SLIDE_W,
            height=Inches(6.6),
        )

    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.55), Inches(8.5), Inches(2.2))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # 1. Eyebrow
    p0 = tf.paragraphs[0]
    p0.text = "CPAG REVIEW PACK"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = BRAND_BLUE
    p0.space_after = Pt(8)

    # 2. Title
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(0x12, 0x26, 0x3F)
    p1.space_after = Pt(8)

    # 3. Subtitle
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(0x2B, 0x4F, 0x6E)
    p2.space_after = Pt(12)

    # Accent bar
    line_shape = slide.shapes.add_shape(
        1, Inches(0.65), Inches(2.95), Inches(1.2), Inches(0.035)
    )
    line_shape.fill.solid()
    line_shape.fill.fore_color.rgb = BRAND_BLUE
    line_shape.line.color.rgb = BRAND_BLUE

    # Stamp
    if stamp:
        box_stamp = slide.shapes.add_textbox(Inches(0.65), Inches(3.1), Inches(9.5), Inches(0.5))
        tf_s = box_stamp.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = stamp
        p_s.font.size = Pt(10)
        p_s.font.color.rgb = RGBColor(0x3D, 0x5A, 0x73)


def _thank_you(prs) -> None:
    slide = prs.slides.add_slide(_layout(prs, "4_Section Header"))
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 0:
            shape.text_frame.text = "Thank You"
    _drop_empty_placeholders(slide)


# ── Populated slides ───────────────────────────────────────────────────────
def slide_capacity(prs, projects: List[Dict[str, Any]], meta) -> None:
    slide = _content_slide(prs, "Battery Energy Storage Systems",
                           "Projects : FY 2026-27")
    rows = [["PSS", "Power", "Energy", "Dispatchable Energy", "Containers"]]
    for p in projects:
        rows.append([p["pss"], f"{_num(p['powerMw'])} MW",
                     f"{_num(p['energyMwh'])} MWh",
                     f"{_num(p['dispatchableMwh'])} MWh",
                     _num(p["containers"])])
    rows.append(["Total", f"{_num(meta['powerMw'])} MW",
                 f"{_num(meta['energyMwh'])} MWh",
                 f"{_num(meta['dispatchableMwh'])} MWh",
                 _num(meta["containers"])])
    _table(slide, rows, MARGIN + Inches(2.0), CONTENT_TOP + Inches(0.4),
           Inches(8.3), col_w=[1.6, 1.6, 1.7, 2.2, 1.4], size=12, row_h=0.42)
    _footnote(slide,
              f"Power, energy, dispatchable energy and container counts are "
              f"declared by the {meta['declaredSource']} and have no system of "
              f"record.")


def slide_salient(prs, projects: List[Dict[str, Any]], meta, commercial) -> None:
    slide = _content_slide(prs, "Salient Features", "All BESS projects")
    rows = [["PSS", "SPV", "Plot", "Battery OEM", "Containers",
             "Land (acres)", "Approvals", "Ordered Rs Cr"]]
    for p in projects:
        rows.append([p["pss"], p["spv"], p.get("plot", "-"), p["batteryOem"],
                     _num(p["containers"]), _num(p.get("landAcres"), 2),
                     f"{p['approvalsDone']}/{p['approvalsTotal']}",
                     _num(p["orderCr"], 1) if p["sapAvailable"] else "civil only"])
    _table(slide, rows, MARGIN, CONTENT_TOP, BODY_W,
           col_w=[1.2, 1.2, 1.0, 1.6, 1.2, 1.3, 1.2, 1.6], size=10, row_h=0.36)
    _text(slide, MARGIN, CONTENT_TOP + Inches(3.1), BODY_W, Inches(1.4),
          f"Project capacity: {_num(meta['powerMw'])} MW / "
          f"{_num(meta['energyMwh'])} MWh (dispatchable "
          f"{_num(meta['dispatchableMwh'])} MWh)\n"
          f"Project location: Khavda RE Park, Gujarat\n"
          f"Order book: Rs {_num(commercial['orderCr'], 1)} Cr across "
          f"{_num(commercial['poCount'])} purchase orders, "
          f"Rs {_num(commercial['deliveredCr'], 1)} Cr delivered",
          size=12)
    _footnote(slide,
              "Capex NFA, PPA effective date, tariff and land lease terms are "
              "commercial facts held outside the connected systems.")


def _bucket_remark(b: Dict[str, Any], p: Dict[str, Any]) -> str:
    key = b.get("key")
    within = b.get("withinBucketPct")
    if key == "statutory":
        done = p.get("approvalsDone") or 0
        tot = p.get("approvalsTotal") or 0
        return f"{done} of {tot} approvals complete; {tot - done} outstanding."
    elif key == "engineering":
        e = p.get("engineering") or {}
        tot = e.get("total") or 0
        done = e.get("completed") or 0
        return f"{done} of {tot} engineering activities complete." if tot else "Engineering activities complete."
    elif key in ("procurement", "ordering"):
        packages = p.get("packages") or []
        pkgs_count = len(packages)
        recd = sum(
            1 for pk in packages
            if any("receipt at site" in (m.get("name") or "").lower() and m.get("actualFinish")
                   for m in pk.get("milestones") or [])
        )
        delivered_cr = _num(p.get("deliveredCr"), 1)
        order_cr = _num(p.get("orderCr"), 1)
        return (f"{recd} of {pkgs_count} packages have receipts at site; "
                f"₹{delivered_cr} Cr of ₹{order_cr} Cr delivered.")
    else:
        q = p.get("quality") or {}
        open_nc = _num(q.get("openNc"))
        rfi_tot = _num(q.get("rfiTotal"))
        earned_scope = _pct(within, 0)
        return f"{earned_scope} of construction scope earned; {open_nc} non-conformances open against {rfi_tot} inspections."


def slide_scurve(prs, pss: str, spv: str, curve, buckets,
                 plan_total: float, actual_total: float,
                 project: Optional[Dict[str, Any]] = None) -> None:
    slide = _content_slide(prs, f"Physical Progress : {pss} S-Curve", spv)
    project = project or {}

    chart_top = CONTENT_TOP
    chart_h = Inches(1.95)

    categories = [_month(c["month"]) for c in curve] if curve else []

    # 1. Full-width combo chart: Monthly bars + Cumulative lines
    if curve and categories:
        cd = CategoryChartData()
        cd.categories = categories
        cd.add_series("Monthly plan", [c.get("planMonthPct") or 0.0 for c in curve])
        cd.add_series("Monthly actual", [c.get("actualMonthPct") or 0.0 if c.get("actualMonthPct") is not None else 0.0 for c in curve])
        cd.add_series("Cumulative plan", [c.get("planCumPct") or 0.0 for c in curve])
        cd.add_series("Cumulative actual", [c.get("actualCumPct") if c.get("actualCumPct") is not None else 0.0 for c in curve])

        shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, chart_top, BODY_W, chart_h, cd
        )
        chart = shape.chart
        chart.has_title = False
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.font.size = Pt(8.5)

        plot_area = chart._element.plotArea
        bar_chart = plot_area.find(qn("c:barChart"))
        if bar_chart is not None:
            series_list = list(bar_chart.findall(qn("c:ser")))
            ax_ids = bar_chart.findall(qn("c:axId"))
            if len(ax_ids) >= 2 and len(series_list) >= 4:
                cat_ax_id = ax_ids[0].get("val")
                val_ax_id = ax_ids[1].get("val")

                line_chart = parse_xml(f'''
                    <c:lineChart {nsdecls("c")}>
                        <c:grouping val="standard"/>
                        <c:axId val="{cat_ax_id}"/>
                        <c:axId val="{val_ax_id}"/>
                    </c:lineChart>
                ''')
                first_ax = line_chart.find(qn("c:axId"))
                for s in series_list[2:]:
                    bar_chart.remove(s)
                    line_chart.insert(line_chart.index(first_ax), s)

                plot_area.insert(plot_area.index(bar_chart) + 1, line_chart)

            plot = chart.plots[0]
            if len(plot.series) >= 2:
                plot.series[0].format.fill.solid()
                plot.series[0].format.fill.fore_color.rgb = BRAND_BLUE
                plot.series[1].format.fill.solid()
                plot.series[1].format.fill.fore_color.rgb = BRAND_PURPLE

    # 2. Monthly grid table across full width
    if curve and categories:
        grid_headers = [""] + categories
        row_m_plan = ["Monthly plan"] + [
            f"{c.get('planMonthPct'):.2f}%" if c.get("planMonthPct") is not None else "—"
            for c in curve
        ]
        row_m_act = ["Monthly actual"] + [
            f"{c.get('actualMonthPct'):.2f}%" if c.get("actualMonthPct") is not None else "—"
            for c in curve
        ]
        row_c_plan = ["Cumulative plan"] + [
            f"{c.get('planCumPct'):.2f}%" if c.get("planCumPct") is not None else "—"
            for c in curve
        ]
        row_c_act = ["Cumulative actual"] + [
            f"{c.get('actualCumPct'):.2f}%" if c.get("actualCumPct") is not None else "—"
            for c in curve
        ]

        grid_rows = [grid_headers, row_m_plan, row_m_act, row_c_plan, row_c_act]
        n_cols = len(grid_headers)
        label_w = 1.6
        each_w = (12.29 - label_w) / max(1, n_cols - 1)
        col_w = [label_w] + [each_w] * (n_cols - 1)

        grid_top = chart_top + chart_h + Inches(0.08)
        _table(slide, grid_rows, MARGIN, grid_top, BODY_W,
               col_w=col_w, size=7.5, row_h=0.18, head_h=0.20)
        v_top = grid_top + Inches(0.20 + 0.18 * 4 + 0.10)
    else:
        v_top = chart_top + Inches(2.2)

    # 3. Variance & Remarks table across full width (6 columns), with the
    # pack's own merged "FTM" header spanning the Plan/Actual pair.
    v_top_groups = [("", 1), ("", 1), ("FTM", 2), ("", 1), ("", 1)]
    v_head = ["Activity", "Wtg.", "Plan", "Actual",
              "Variance in Plan Vs Actual",
              "Variance Remark & Mitigation / Action Plan"]
    v_rows: List[List[Any]] = []
    tints = {}
    for i, b in enumerate(buckets):
        v_rows.append([
            b["label"].split(",")[0],
            _pct(b["weightPct"], 0),
            _pct(b["planToDatePct"], 0),
            _pct(b["earnedPct"], 0),
            "-" if b["variancePct"] is None else f"{b['variancePct']:+.1f}",
            _bucket_remark(b, project)
        ])
        tints[i] = _variance_tint(b["variancePct"])

    v_rows.append([
        "Total", "100%", _pct(plan_total, 0), _pct(actual_total, 0),
        f"{actual_total - plan_total:+.1f}",
        "P6 weightage units, phased on baseline and actual finish dates. The pack reports construction on installed physical quantity, tracked outside Akasha, so that row may differ."
    ])
    tints[len(v_rows) - 1] = _variance_tint(actual_total - plan_total)

    _table_grouped(slide, v_top_groups, v_head, v_rows, MARGIN, v_top, BODY_W,
                   col_w=[1.8, 0.6, 0.6, 0.6, 1.2, 7.49],
                   size=8, row_h=0.26, head_h=0.24, tints=tints, tint_col=4)

    _footnote(slide,
              "P6 weightage units - planned against actual, phased on baseline "
              "and actual finish dates. The pack reports the construction bucket "
              "on installed physical quantity, which is tracked outside Akasha, "
              "so that line may differ from the circulated pack.")


def _milestone_of(pkg, *needles):
    for m in pkg["milestones"]:
        if any(n in m["name"].lower() for n in needles):
            return m
    return None


def _receipt_milestones(pkg) -> List[Dict[str, Any]]:
    return sorted(
        (m for m in pkg["milestones"]
         if "receipt at site" in m["name"].lower()),
        key=lambda m: m.get("baselineFinish") or "",
    )


def _actual_or_forecast(m) -> str:
    """Actual where the milestone landed; P6's own forecast finish where it
    hasn't yet - never blank just because the work is still ahead of today.
    The asterisk marks a forecast so it reads as a projection, not a fact."""
    if not m:
        return "-"
    if m.get("actualFinish"):
        return _date(m["actualFinish"])
    if m.get("forecastFinish"):
        return f"{_date(m['forecastFinish'])}*"
    return "-"


def slide_procurement(prs, pss: str, packages, sap, sap_note,
                      project: Optional[Dict[str, Any]] = None) -> None:
    """Seventeen columns with the pack's own two-row merged header - Sr. No.
    through Remarks, "Expected Delivery at Site" over Start/Finish,
    "Manufacturing Status" over QAP Acceptance/MC Date, and "Forecast
    Delivery Schedule (Qty)" over the next two reporting months.
    """
    slide = _content_slide(
        prs,
        f"Procurement & Monthly Rolling Plan for Supply of Equipment - {pss}",
        "P6 ordering milestones joined to SAP purchase orders")
    by_material = {r["material"]: r for r in (sap or [])}
    fc_months = _next_months(_last_reported_month(project), 2)
    fc_labels = [_month(m) for m in fc_months]

    top_groups = [
        ("", 1), ("", 1), ("", 1), ("", 1), ("", 1), ("", 1), ("", 1),
        ("", 1), ("", 1),
        ("Expected Delivery at Site", 2),
        ("Manufacturing Status", 2),
        ("", 1),
        ("Forecast Delivery Schedule (Qty)", 2),
        ("", 1),
    ]
    head = ["Sr. No.", "Packages", "Manufacturer", "UOM", "Scope",
            "Ordering Completed", "Balance Ordering", "PO Number", "PO Date",
            "Start", "Finish", "QAP Acceptance Date", "MC Date",
            "Delivered At Site", fc_labels[0], fc_labels[1], "Remarks"]

    rows = []
    for i, pkg in enumerate(packages, start=1):
        order = _milestone_of(pkg, "placement of the order")
        mc = _milestone_of(pkg, "manufacturing clearance", "ntp")
        receipts = _receipt_milestones(pkg)
        done = [m for m in receipts if m["actualFinish"]]
        landed = sorted((m["actualFinish"] for m in done))
        slips = [m["slipDays"] for m in pkg["milestones"]
                 if m["slipDays"] is not None]
        worst = max(slips) if slips else None
        sap_row = by_material.get(pkg.get("sapMaterial") or "")
        scope = pkg.get("scopeQty")
        ordered = sap_row.get("orderQtyRaw") if sap_row else None
        # Balance is only meaningful when scope and ordered share a unit; SAP
        # quantities mix units across multi-line POs for every material
        # except containers, so it is otherwise left blank rather than guessed.
        balance = "-"
        if (scope is not None and ordered is not None
                and sap_row and sap_row.get("material") == "BESS Containers"):
            balance = _num(max(scope - ordered, 0))
        remark_bits = []
        if receipts:
            remark_bits.append(f"{len(done)} of {len(receipts)} lots received")
        if landed:
            remark_bits.append(f"last {_date(landed[-1])}")
        if worst is not None and worst > 0:
            remark_bits.append(f"worst slip +{worst}d")
        # Expected Delivery at Site: first and last lot, actual where it has
        # landed, P6's own forecast where it hasn't - the same pattern as PO
        # Date and MC Date below, so a reader sees one convention throughout.
        delivery_start = _actual_or_forecast(receipts[0]) if receipts else "-"
        delivery_finish = _actual_or_forecast(receipts[-1]) if receipts else "-"
        rows.append([
            str(i),
            pkg.get("packageBase") or pkg.get("package"),
            (sap_row["vendor"] if sap_row else "-") or "-",
            pkg.get("uom") or ("Nos." if scope else "-"),
            _num(scope) if scope else "-",
            _num(ordered) if ordered is not None else "-",
            balance,
            (sap_row["poNumbers"] if sap_row else "-") or "-",
            _actual_or_forecast(order),
            delivery_start, delivery_finish,
            "-", _actual_or_forecast(mc),
            _num(sap_row["deliveredQtyRaw"]) if sap_row and sap_row.get("deliveredQtyRaw") else "-",
            "-", "-",
            "; ".join(remark_bits) if remark_bits else "-",
        ])
    # The reference tints this table by column, not by row: grey over the
    # nine ordering columns (Sr. No. through PO Date), peach over the eight
    # delivery-tracking columns (Expected Delivery through Remarks).
    col_tint = [GREY_COL] * 9 + [PEACH_COL] * 8
    _table_grouped(
        slide, top_groups, head, rows, MARGIN, CONTENT_TOP, BODY_W,
        col_w=[0.35, 1.3, 1.25, 0.4, 0.45, 0.65, 0.65, 1.05, 0.6, 0.55, 0.55,
               0.65, 0.6, 0.6, 0.42, 0.42, 1.4],
        size=7.2, row_h=0.27, head_h=0.24, col_tint=col_tint)
    base_note = sap_note or (
        "SAP holds no PO date or delivery date for these projects, so dates "
        "are P6 ordering milestones; manufacturer, PO number and delivered "
        "quantity are from SAP. QAP acceptance date and forecast delivery "
        "quantities have no source in any connected system.")
    _footnote(slide, base_note + " * marks a forecast date - P6's current "
              "schedule finish for a milestone that has not happened yet.")


def _approval_remark(a: Dict[str, Any]) -> str:
    if a["status"] == "Completed":
        slip = a.get("slipDays")
        if slip is None or slip <= 0:
            return f"Completed {_date(a['actualFinish'])}."
        return f"Completed {_date(a['actualFinish'])} - {slip}d behind baseline."
    if a["status"] == "In Progress":
        return "In progress."
    if a.get("forecastFinish"):
        return f"Forecast {_date(a['forecastFinish'])}."
    return "-"


def slide_approvals(prs, pss: str, items, done: int, total: int,
                    start: int, end: int, page: str) -> None:
    """Nine columns, exactly as the pack has them: SN through Remarks, no
    separate slip column - slip is folded into the Remarks text instead."""
    slide = _content_slide(prs, f"Statutory and Other Approvals ({pss})",
                           f"{done} of {total} complete · {page}")
    rows = [["SN", "Activity Name", "Responsible", "Approval Authority",
             "Baseline Start", "Baseline Finish", "Actual/ Forecast Start",
             "Actual/ Forecast Finish", "Remarks"]]
    for i, a in enumerate(items[start:end], start=start + 1):
        rows.append([
            str(i), a["name"], "-", "-",
            _date(a["baselineStart"]), _date(a["baselineFinish"]),
            _date(a["actualStart"]),
            _date(a["actualFinish"] or a["forecastFinish"]),
            _approval_remark(a),
        ])
    _table(slide, rows, MARGIN, CONTENT_TOP, BODY_W,
           col_w=[0.35, 2.6, 1.0, 1.25, 1.0, 1.0, 1.1, 1.15, 3.84],
           size=8.5, row_h=0.30)
    _footnote(slide,
              "Responsible party and approval authority are narrative columns "
              "in the pack with no system of record, so they are blank.")


def slide_ordering_status(prs, pss_group: str, kind: str, rows_body,
                          page: str) -> None:
    """Eleven columns, exactly as the pack has them, including the "Verified
    by Engineering Consultant" column the pack always prints as NA."""
    slide = _content_slide(prs, f"{kind} ordering Status ({pss_group})", page)
    header = ["SN", "Equipment/ Package", "Qty / Scope",
              "Final / Revised Specifications Issue Date",
              "Verified by Engineering Consultant", "TBER Date",
              "Qualified Vendors", "Receipt of Offer",
              "Commercial NFA Submission", "Release of PO / SO", "Update"]
    _table(slide, [header] + (rows_body or []), MARGIN, CONTENT_TOP, BODY_W,
           col_w=[0.35, 1.9, 0.85, 1.25, 1.1, 0.95, 1.15, 0.95, 1.1, 1.15, 1.54],
           size=9.5, row_h=0.40, head_h=0.34)
    _footnote(slide,
              "Specification issue date, TBER date, qualified vendor list, "
              "receipt of offer and commercial NFA submission are not held in "
              "any connected system and are left blank. Release of PO/SO is "
              "the P6 order-placement milestone; * marks a forecast date "
              "where the order has not yet been placed. Vendor and value "
              "come from SAP.")


def slide_mandays(prs, series, planned: int, earned: int, title_suffix: str,
                  posted: Optional[int] = None) -> None:
    """Chart and summary on top, then the per-month grid underneath - the
    same layout as the S-curve and Financial slides, so a reader can take a
    number off the table without reading the chart."""
    slide = _content_slide(
        prs, f"Mandays Deployment Plan vs Actual - {title_suffix}",
        "Scope (planned Labor units) against earned mandays")
    chart_h = Inches(2.35)
    if series:
        _bar_chart(slide, [_month(m["month"]) for m in series],
                   [("Planned (month)", [m.get("planMonth", 0) for m in series]),
                    ("Earned (month)", [m.get("earnedMonth", 0) for m in series])],
                   MARGIN, CONTENT_TOP, Inches(7.15), chart_h)
    rows = [["Measure", "Mandays"],
            ["Scope (planned Labor units)", _num(planned)],
            ["Earned (scope x % complete)", _num(earned)],
            ["Remaining", _num(planned - earned)],
            ["Earned %", _pct((earned / planned * 100) if planned else None)]]
    if posted is not None:
        rows.append(["Posted actual units", _num(posted)])
    _table(slide, rows, MARGIN + Inches(7.35), CONTENT_TOP, Inches(4.94),
           col_w=[3.0, 1.94], size=10, row_h=0.42)

    if series:
        months = [_month(m["month"]) for m in series]
        grid_rows = [
            [""] + months,
            ["Planned (month)"] + [_num(m.get("planMonth", 0)) for m in series],
            ["Earned (month)"] + [_num(m.get("earnedMonth", 0)) for m in series],
            ["Planned cum."] + [_num(m.get("planCum", 0)) for m in series],
            ["Earned cum."] + [_num(m.get("earnedCum", 0)) for m in series],
        ]
        label_w = 1.6
        each_w = (12.29 - label_w) / max(1, len(months))
        _table(slide, grid_rows, MARGIN, CONTENT_TOP + chart_h + Inches(0.1),
               BODY_W, col_w=[label_w] + [each_w] * len(months),
               size=7.5, row_h=0.24)

    _footnote(slide,
              "Scope is planned Labor units (mandays) on the schedule. Earned "
              "= scope x activity percent complete - the same percent that "
              "drives physical progress - because actual labour units are "
              "not posted in P6 (1-4 non-zero rows out of 1,300-3,300 per "
              "project). This shows mandays credited for work done, not "
              "workmen deployed, and cannot reveal a productivity gap.")


def slide_manpower_curve(prs, series, title_suffix: str) -> None:
    slide = _content_slide(
        prs, f"Manpower Deployment Plan vs Actual - {title_suffix}",
        "Cumulative mandays")
    if series:
        _line_chart(slide, [_month(m["month"]) for m in series],
                    [("Planned cumulative", [m.get("planCum", 0) for m in series]),
                     ("Earned cumulative", [m.get("earnedCum", 0) for m in series])],
                    MARGIN, CONTENT_TOP, BODY_W, Inches(4.5),
                    number_format='#,##0')
    _footnote(slide,
              "Average headcount can be read as mandays divided by working days "
              "in the month. Contractor-level deployment is not held in any "
              "connected system.")


def slide_financial(prs, financial: Dict[str, Any], label: str) -> None:
    """The pack's own Financial S-Curve layout: one combo chart, then a
    four-row month grid - Budgeted cumulative / Monthly Budgeted / Actual
    Cumulative / Monthly Actual - exactly as slide 45 has it, rather than the
    per-PSS ordered/delivered table this used to substitute.
    """
    slide = _content_slide(prs, "Financial S Curve", label)
    series = financial.get("series") or []
    months = [_month(pt["month"]) for pt in series]

    if series:
        cd = CategoryChartData()
        cd.categories = months
        cd.add_series("Monthly Budgeted", [pt["budgetedMonthCr"] for pt in series])
        cd.add_series("Monthly Actual", [pt["actualMonthCr"] for pt in series])
        cd.add_series("Budgeted- cumulative", [pt["budgetedCumCr"] for pt in series])
        cd.add_series("Actual Cumulative", [pt["actualCumCr"] for pt in series])
        shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, CONTENT_TOP,
            BODY_W, Inches(3.1), cd)
        chart = shape.chart
        chart.has_title = False
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.font.size = Pt(8.5)
        plot_area = chart._element.plotArea
        bar_chart = plot_area.find(qn("c:barChart"))
        if bar_chart is not None:
            series_list = list(bar_chart.findall(qn("c:ser")))
            ax_ids = bar_chart.findall(qn("c:axId"))
            if len(ax_ids) >= 2 and len(series_list) >= 4:
                cat_ax_id, val_ax_id = ax_ids[0].get("val"), ax_ids[1].get("val")
                line_chart = parse_xml(f'''
                    <c:lineChart {nsdecls("c")}>
                        <c:grouping val="standard"/>
                        <c:axId val="{cat_ax_id}"/>
                        <c:axId val="{val_ax_id}"/>
                    </c:lineChart>
                ''')
                first_ax = line_chart.find(qn("c:axId"))
                for s in series_list[2:]:
                    bar_chart.remove(s)
                    line_chart.insert(line_chart.index(first_ax), s)
                plot_area.insert(plot_area.index(bar_chart) + 1, line_chart)
            plot = chart.plots[0]
            if len(plot.series) >= 2:
                plot.series[0].format.fill.solid()
                plot.series[0].format.fill.fore_color.rgb = BRAND_BLUE
                plot.series[1].format.fill.solid()
                plot.series[1].format.fill.fore_color.rgb = BRAND_PURPLE

    # The month grid beneath the chart, in the pack's own row order.
    grid_top = CONTENT_TOP + Inches(3.2)
    header = [""] + months
    rows = [
        header,
        ["Budgeted- cumulative"] + [_num(pt["budgetedCumCr"], 0) for pt in series],
        ["Monthly Budgeted"] + [_num(pt["budgetedMonthCr"], 0) for pt in series],
        ["Actual Cumulative"] + [_num(pt["actualCumCr"], 0) for pt in series],
        ["Monthly Actual"] + [_num(pt["actualMonthCr"], 0) for pt in series],
    ]
    label_w = 1.7
    each_w = (12.29 - label_w) / max(1, len(months))
    _table(slide, rows, MARGIN, grid_top, BODY_W,
           col_w=[label_w] + [each_w] * len(months), size=7.5, row_h=0.24)
    _footnote(slide,
              "Budgeted is derived, not measured: the Rs 16,358 Cr approved "
              "NFA (NFA_R1, 04-Apr-26) spread across each project's own P6 "
              "baseline schedule, weighted by dispatchable MWh share - a "
              "schedule projection onto a real total, not a cost-loaded plan. "
              "Actual is SAP goods-receipt value and is currently zero "
              "portfolio-wide because mt_materialdocument, the only source "
              "with a posting date to phase it by month, holds no rows at "
              "this sync.")


def slide_contractors(prs, pss: str, items, note: str) -> None:
    site = [c for c in items if c.get("kind") == "site"]
    slide = _content_slide(prs, f"Contractor Status ({pss})",
                           "Contracted scope and quality record")
    rows = [["Contractor", "POs", "Contracted Rs Cr", "Delivered Rs Cr",
             "Delivered %", "RFIs", "NCs", "NC rate"]]
    for c in site[:15]:
        rows.append([c["vendor"], str(c["poCount"]), _num(c["orderCr"], 2),
                     _num(c["deliveredCr"], 2), _pct(c["deliveredPct"], 0),
                     str(c["rfiCount"] or "-"), str(c["ncCount"] or "-"),
                     _pct(c["ncRatePct"], 1)])
    _table(slide, rows, MARGIN, CONTENT_TOP, BODY_W,
           col_w=[4.0, 0.7, 1.6, 1.6, 1.2, 0.9, 0.8, 1.0], size=9, row_h=0.30)
    _footnote(slide, note)

def slide_engineering(prs, pss: str, spv: str, eng: Dict[str, Any]) -> None:
    """Engineering activity progress. The pack prints a document master log we
    do not ingest, so the slide reports what P6 holds and names the difference."""
    slide = _content_slide(prs, f"Engineering Progress : {pss}", spv)
    m = eng.get("monthly") or []
    if m:
        _bar_chart(slide, [_month(x["month"]) for x in m],
                   [("Plan (month)", [x["planMonth"] for x in m]),
                    ("Actual (month)", [x["actualMonth"] for x in m])],
                   MARGIN, CONTENT_TOP, Inches(6.9), Inches(4.2))
    rows = [["Description", "Activities", "Completed", "Complete %"]]
    for g in eng.get("groups") or []:
        rows.append([g["name"], _num(g["total"]), _num(g["completed"]),
                     _pct(g["pct"], 0)])
    total, done = eng.get("total") or 0, eng.get("completed") or 0
    rows.append(["Total", _num(total), _num(done),
                 _pct((done / total * 100) if total else None, 0)])
    _table(slide, rows, MARGIN + Inches(7.15), CONTENT_TOP, Inches(5.14),
           col_w=[2.4, 1.0, 1.0, 0.9], size=10.5, row_h=0.36)
    _footnote(slide, eng.get("basis") or "")


def slide_commissioning(prs, pss: str, spv: str, comm: Dict[str, Any]) -> None:
    """Commissioning and trial-run milestones by month."""
    slide = _content_slide(prs, f"Commissioning Plan : {pss}", spv)
    m = comm.get("monthly") or []
    if m:
        _bar_chart(slide, [_month(x["month"]) for x in m],
                   [("Milestones due", [x["milestones"] for x in m])],
                   MARGIN, CONTENT_TOP, Inches(5.4), Inches(4.2),
                   number_format="0")
    rows = [["Milestone", "Status", "Baseline", "Forecast / actual", "Slip"]]
    for it in (comm.get("items") or [])[:13]:
        rows.append([it["name"], it["status"], _date(it["baselineFinish"]),
                     _date(it["actualFinish"] or it["forecastFinish"]),
                     "-" if it["slipDays"] is None else f"{it['slipDays']:+d}d"])
    _table(slide, rows, MARGIN + Inches(5.65), CONTENT_TOP, Inches(6.64),
           col_w=[2.5, 1.1, 1.0, 1.3, 0.74], size=9.5, row_h=0.31)
    _footnote(slide, comm.get("basis") or "")


def slide_quality(prs, pss: str, qual: Dict[str, Any]) -> None:
    """Pulse inspection and non-conformance record.

    Not a slide in the circulated pack. It is included because it is measured
    per project and is the nearest evidence for the construction variance the
    pack reports.
    """
    slide = _content_slide(prs, f"Quality Record : {pss}",
                           "Inspection requests and non-conformances from Pulse")
    m = qual.get("monthly") or []
    if m:
        _bar_chart(slide, [_month(x["month"]) for x in m],
                   [("Inspections raised", [x["rfi"] for x in m]),
                    ("Non-conformances", [x["nc"] for x in m])],
                   MARGIN, CONTENT_TOP, Inches(6.4), Inches(4.2),
                   number_format="0")
    rows = [["Measure", "Value"],
            ["Inspections raised", _num(qual.get("rfiTotal"))],
            ["Non-conformances", _num(qual.get("ncTotal"))],
            ["Open NCs", _num(qual.get("openNc"))],
            ["NC rate", _pct(qual.get("ncRatePct"), 1)]]
    for r in (qual.get("ncByStatus") or [])[:5]:
        rows.append([f"NC — {r['status']}", _num(r["count"])])
    for r in (qual.get("rfiByPackage") or [])[:4]:
        rows.append([f"Inspections — {r['package']}", _num(r["count"])])
    _table(slide, rows, MARGIN + Inches(6.65), CONTENT_TOP, Inches(5.64),
           col_w=[3.4, 2.24], size=10, row_h=0.32)
    _footnote(slide,
              f"Pulse records for {qual.get('project')}. Not a slide in the "
              "circulated pack - included because it is measured per project "
              "and is the nearest evidence for the construction variance.")

def slide_engineering_all(prs, projects: List[Dict[str, Any]], label: str) -> None:
    """Engineering across every project in scope — one slide, as the pack has it."""
    slide = _content_slide(prs, "Engineering Progress", label)
    rows = [["PSS", "Basic", "Detailed", "Technical Spec.", "Total activities",
             "Completed", "Complete %"]]
    for p in projects:
        eng = p.get("engineering") or {}
        by = {g["name"]: g for g in eng.get("groups") or []}

        def cell(name: str) -> str:
            g = by.get(name)
            return f"{g['completed']}/{g['total']}" if g else "-"

        total, done = eng.get("total") or 0, eng.get("completed") or 0
        rows.append([p["pss"], cell("Basic Engineering"),
                     cell("Detailed Engineering"), cell("Technical Sepcification"),
                     _num(total), _num(done),
                     _pct((done / total * 100) if total else None, 0)])
    _table(slide, rows, MARGIN, CONTENT_TOP, BODY_W,
           col_w=[1.4, 1.4, 1.6, 1.8, 1.8, 1.4, 1.4], size=11, row_h=0.4)
    basis = next((p.get("engineering", {}).get("basis") for p in projects
                  if (p.get("engineering") or {}).get("basis")), "")
    _footnote(slide, basis)


def slide_commissioning_all(prs, projects: List[Dict[str, Any]], label: str) -> None:
    """Commissioning milestones by month across every project — one slide."""
    slide = _content_slide(prs, "Commissioning Plan", label)
    months: List[str] = []
    for p in projects:
        for m in (p.get("commissioning") or {}).get("monthly") or []:
            if m["month"] not in months:
                months.append(m["month"])
    months.sort()

    rows = [["PSS"] + [_month(m) for m in months] + ["Total"]]
    totals = [0] * len(months)
    for p in projects:
        by = {m["month"]: m["milestones"]
              for m in (p.get("commissioning") or {}).get("monthly") or []}
        line = [p["pss"]]
        n = 0
        for i, m in enumerate(months):
            v = by.get(m, 0)
            totals[i] += v
            n += v
            line.append(str(v) if v else "-")
        line.append(str(n))
        rows.append(line)
    rows.append(["Total"] + [str(t) for t in totals] + [str(sum(totals))])

    widths = [1.4] + [1.0] * len(months) + [1.0]
    _table(slide, rows, MARGIN, CONTENT_TOP, BODY_W,
           col_w=widths, size=11, row_h=0.4)
    basis = next((p.get("commissioning", {}).get("basis") for p in projects
                  if (p.get("commissioning") or {}).get("basis")), "")
    _footnote(slide, basis + " Counts are commissioning and trial-run "
              "milestones falling in each month, not MWh released.")


def slide_quality_all(prs, projects: List[Dict[str, Any]], label: str) -> None:
    """Quality across every project in scope — one slide."""
    slide = _content_slide(prs, "Quality Record",
                           "Inspection requests and non-conformances from Pulse")
    rows = [["PSS", "Inspections raised", "Non-conformances", "Open NCs",
             "NC rate", "Civil RFIs", "Electrical RFIs"]]
    for p in projects:
        q = p.get("quality") or {}
        by_pkg = {r["package"]: r["count"] for r in q.get("rfiByPackage") or []}
        rows.append([p["pss"], _num(q.get("rfiTotal")), _num(q.get("ncTotal")),
                     _num(q.get("openNc")), _pct(q.get("ncRatePct"), 1),
                     _num(by_pkg.get("Civil", 0)),
                     _num(by_pkg.get("Electrical", 0))])
    _table(slide, rows, MARGIN, CONTENT_TOP, BODY_W,
           col_w=[1.4, 2.0, 2.0, 1.4, 1.4, 1.5, 1.7], size=11, row_h=0.4)
    _footnote(slide,
              "Not a slide in the circulated pack - included because it is "
              "measured per project and is the nearest evidence for the "
              "construction variance the pack reports.")

# The pack colours an actual bar by how much of plan it reached.
PLAN_BLUE = RGBColor(0x2E, 0x9B, 0xD6)
ACH_GREEN = RGBColor(0x00, 0xA6, 0x50)
ACH_YELLOW = RGBColor(0xFF, 0xD1, 0x1A)
ACH_RED = RGBColor(0xE8, 0x11, 0x2D)


def _achievement(pct) -> RGBColor:
    if pct is None:
        return MUTED
    if pct >= 80:
        return ACH_GREEN
    if pct >= 40:
        return ACH_YELLOW
    return ACH_RED


def slide_civil(prs, pss: str, spv: str, con: Dict[str, Any]) -> None:
    """Civil construction progress by element, from P6 Material quantities."""
    slide = _content_slide(
        prs, f"Civil Construction Progress (Plan vs Actual) ({pss})", spv)
    elements = [e for e in (con.get("elements") or [])
                if e["element"] != "General"][:10]
    rows = [["Element", "Stages", "Planned", "Actual", "Progress"]]
    tints = {}
    for i, e in enumerate(elements, start=1):
        rows.append([e["element"], str(e["stageCount"]), _num(e["planned"]),
                     _num(e["actual"]), _pct(e["pct"], 0)])
        tints[i] = _achievement(e["pct"])
    _table(slide, rows, MARGIN, CONTENT_TOP, Inches(5.9),
           col_w=[1.8, 0.9, 1.2, 1.2, 0.9], size=10, row_h=0.33, tints=tints)

    if elements:
        chart = _bar_chart(
            slide, [e["element"] for e in elements],
            [("Plan", [100 for _ in elements]),
             ("Actual", [e["pct"] or 0 for e in elements])],
            MARGIN + Inches(6.15), CONTENT_TOP, Inches(6.14), Inches(4.1),
            number_format='0"%"')
        plots = chart.plots[0].series
        plots[0].format.fill.solid()
        plots[0].format.fill.fore_color.rgb = PLAN_BLUE
        for idx, pt in enumerate(plots[1].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _achievement(elements[idx]["pct"])
    _footnote(slide, con.get("basis") or "")


def slide_electrical(prs, pss: str, spv: str, elec: Dict[str, Any]) -> None:
    """Electrical progress on the pack's own element labels.

    The pack prints Scope and Plan as separate columns even though they hold
    the same figure here - both are kept so the column count matches.
    """
    slide = _content_slide(prs, f"Electrical Plan Vs Actual Summary {pss}", spv)
    items = elec.get("items") or []
    rows = [["Element", "UoM", "Scope", "Plan", "Actual", "Plan %", "Actual %"]]
    tints = {}
    for i, it in enumerate(items, start=1):
        rows.append([it["element"], it["uom"], _num(it["scope"]),
                     _num(it.get("plan", it["scope"])), _num(it["actual"]),
                     "100%", _pct(it["actualPct"], 0)])
        tints[i] = _achievement(it["actualPct"])
    _table(slide, rows, MARGIN, CONTENT_TOP, Inches(6.15),
           col_w=[1.9, 0.55, 0.95, 0.85, 0.9, 0.75, 0.85], size=9, row_h=0.3,
           tints=tints)

    if items:
        chart = _bar_chart(
            slide, [it["element"] for it in items],
            [("Plan", [100 for _ in items]),
             ("Actual", [it["actualPct"] or 0 for it in items])],
            MARGIN + Inches(6.35), CONTENT_TOP, Inches(5.94), Inches(4.1),
            number_format='0"%"')
        plots = chart.plots[0].series
        plots[0].format.fill.solid()
        plots[0].format.fill.fore_color.rgb = PLAN_BLUE
        for idx, pt in enumerate(plots[1].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _achievement(items[idx]["actualPct"])
    _footnote(slide, elec.get("basis") or "")
