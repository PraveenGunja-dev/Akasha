"""
Split the Khavda monthly CPAG report into one deck per technology.

    python -I scripts/split_khavda_deck.py

Reads  Data/uploads/01. Khavda_CPAG_Energy_Monthly Report_Aug'26.pptx
Writes assets/cpag_khavda_solar_reference.pptx and assets/cpag_khavda_wind_reference.pptx

Each deck keeps the shared opening pages (cover, agenda, executive summary,
commissioning) cut down to its own technology, then only its own section.
Nothing is re-typed: the shared tables and chart are cut from the report's own
cells, so a value shown is a value in the report.
"""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_LABEL_POSITION
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

from services.cpag_pptx import delete_slide, move_slide_to_end, replace_chart_data  # noqa: E402

SOURCE = BACKEND / "Data" / "uploads" / "01. Khavda_CPAG_Energy_Monthly Report_Aug'26.pptx"

# 1-based slide numbers in the source report.
SHARED = [1, 2, 3, 4]
SOLAR = list(range(5, 28)) + [36]   # section divider .. PSS summaries, then tracker & robot status
WIND = list(range(28, 35))          # section divider .. manpower
WIND_MANPOWER = 34                  # combined Solar + Wind chart; cut to Wind below
THANK_YOU = 35

# Agenda paragraphs: 0-1 shared, 2-7 the solar block, 8-12 the wind block.
AGENDA_DROP = {"solar": range(8, 13), "wind": range(2, 8)}

# Commissioning table: rows 2-4 are Budget (Solar, Wind, Total), 6-8 Actual/Fcst.
COMM_BLOCKS = (2, 6)
TECH_ROW = {"solar": 0, "wind": 1}   # offset of the technology's row inside a block
EXEC_ROW = {"solar": 2, "wind": 3}   # executive summary row of each technology
EXEC_TOTAL_ROW = 4
COMM_ROW_HEIGHT = Emu(1_250_000)     # room for the rotated "ACTUAL / FCST" label


def _remove_rows(tbl, rows):
    for r in sorted(rows, reverse=True):
        tbl._tbl.remove(tbl._tbl.tr_lst[r])


def num(s):
    s = s.replace(",", "")
    return float(s) if s else None


def _cut_agenda(slide, tech):
    box = next(sh for sh in slide.shapes if sh.has_text_frame and "Executive Summary" in sh.text_frame.text)
    paras = box.text_frame.paragraphs
    for i in AGENDA_DROP[tech]:
        paras[i]._p.getparent().remove(paras[i]._p)


def _cut_executive_summary(slide, tech):
    tbl = next(sh.table for sh in slide.shapes if getattr(sh, "has_table", False))
    other = EXEC_ROW["wind" if tech == "solar" else "solar"]
    # SN restarts at 1 for the one project left; the Total row would only repeat it.
    tbl.cell(EXEC_ROW[tech], 0).text_frame.paragraphs[0].runs[0].text = "1"
    _remove_rows(tbl, [other, EXEC_TOTAL_ROW])


def _cut_commissioning(slide, tech):
    tbl_shape = next(sh for sh in slide.shapes if getattr(sh, "has_table", False))
    tbl = tbl_shape.table
    trs = tbl._tbl.tr_lst
    keep = TECH_ROW[tech]
    drop = []
    for b in COMM_BLOCKS:
        label = copy.deepcopy(trs[b].tc_lst[0])
        label.attrib.pop("rowSpan", None)
        target = trs[b + keep]
        if keep:  # the label cell sits on the Solar row; carry it onto the kept row
            target.replace(target.tc_lst[0], label)
        else:
            trs[b].replace(trs[b].tc_lst[0], label)
        drop += [b + k for k in range(3) if k != keep]
    _remove_rows(tbl, drop)
    for r in range(len(tbl.rows)):
        if r in (2, 4):  # the two data rows, after the cut
            tbl.rows[r].height = COMM_ROW_HEIGHT
    # Notes under the table follow it up.
    bottom = tbl_shape.top + sum(r.height for r in tbl.rows)
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith(("All nos.", "All the figures")):
            sh.top = bottom + Emu(60_000)

    # The chart is the combined total; redraw it for this technology from the
    # table's own cells (blank month = no bar, not a zero).
    chart = next(sh.chart for sh in slide.shapes if getattr(sh, "has_chart", False))
    cd = CategoryChartData(number_format="#,##0")
    cd.categories = list(chart.plots[0].categories)
    cd.add_series("Budget", [num(tbl.cell(2, c).text_frame.text.strip()) for c in range(2, 14)])
    cd.add_series("Actual / Fcst", [num(tbl.cell(4, c).text_frame.text.strip()) for c in range(2, 14)])
    replace_chart_data(chart, cd)


def _wind_manpower(slide):
    """The report's manpower chart stacks Solar and Wind under overall
    plan/actual/forecast lines. For the wind deck, keep the Wind series only:
    drop the Solar bars and the overall-total overlay (both include solar), and
    label the bars, since the totals that carried the labels are gone."""
    shape = next(sh for sh in slide.shapes if getattr(sh, "has_chart", False))
    chart = shape.chart
    area = chart._chartSpace.find(".//" + qn("c:plotArea"))
    for overlay in area.findall(qn("c:scatterChart")):
        area.remove(overlay)
    bars = area.find(qn("c:barChart"))
    for ser in bars.findall(qn("c:ser")):
        if ser.find(qn("c:tx")).findtext(".//" + qn("c:v")) == "Solar":
            bars.remove(ser)
    labels = chart.plots[0].series[0].data_labels
    for gone in labels._element.findall(qn("c:delete")):  # the series ships with its labels deleted
        labels._element.remove(gone)
    labels.show_value = True
    labels.show_series_name = labels.show_category_name = labels.show_legend_key = False
    labels.number_format, labels.number_format_is_linked = "#,##0", False
    # One series, so stacked and clustered look alike - clustered lets labels sit above the bars.
    bars.find(qn("c:grouping")).set("val", "clustered")
    for overlap in bars.findall(qn("c:overlap")):
        bars.remove(overlap)
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(9)
    chart.has_legend = False  # one series now; the axis already names Plan / Actual / Forecast
    title = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Khavda:"))
    run = title.text_frame.paragraphs[0].runs[0]
    run.text = run.text.replace("Khavda:", "Khavda Wind:", 1)


def build(tech: str, out: Path) -> None:
    prs = Presentation(str(SOURCE))
    slides = list(prs.slides)
    keep = set(SHARED + (SOLAR if tech == "solar" else WIND) + [THANK_YOU])
    for n, s in enumerate(slides, start=1):
        if n not in keep:
            delete_slide(prs, s)
    _cut_agenda(slides[1], tech)
    _cut_executive_summary(slides[2], tech)
    _cut_commissioning(slides[3], tech)
    if tech == "wind":
        _wind_manpower(slides[WIND_MANPOWER - 1])
    move_slide_to_end(prs, slides[THANK_YOU - 1])
    prs.save(str(out))
    print(out.name, len(prs.slides), "slides")


if __name__ == "__main__":
    for tech in ("solar", "wind"):
        build(tech, BACKEND / "assets" / f"cpag_khavda_{tech}_reference.pptx")
