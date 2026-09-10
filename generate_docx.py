from docx import Document

def generate_doc():
    doc = Document()
    doc.add_heading('Akasha Platform Tech Stack', 0)

    doc.add_heading('1. Overview', level=1)
    doc.add_paragraph('Akasha is a cross-platform intelligence system for renewable-energy project portfolios. It acts as an intelligence engine, joining multiple enterprise systems against one canonical project identity.')

    doc.add_heading('2. Architecture & Tech Stack', level=1)

    doc.add_heading('Backend', level=2)
    p = doc.add_paragraph()
    p.add_run('Language: ').bold = True
    p.add_run('Python\n')
    p.add_run('Framework: ').bold = True
    p.add_run('FastAPI\n')
    p.add_run('ORM: ').bold = True
    p.add_run('SQLAlchemy\n')
    p.add_run('Server: ').bold = True
    p.add_run('Uvicorn\n')
    p.add_run('Key Libraries: ').bold = True
    p.add_run('pydantic, celery, redis, openai, langchain-openai, groq, pandas, openpyxl, PyMuPDF, msal, psycopg2-binary, requests')

    doc.add_heading('Database', level=2)
    p = doc.add_paragraph()
    p.add_run('Database Engine: ').bold = True
    p.add_run('PostgreSQL (schema auto-migrated on boot)')

    doc.add_heading('Frontend', level=2)
    p = doc.add_paragraph()
    p.add_run('Framework: ').bold = True
    p.add_run('React 19\n')
    p.add_run('Language: ').bold = True
    p.add_run('TypeScript\n')
    p.add_run('Bundler: ').bold = True
    p.add_run('Vite 8\n')
    p.add_run('Styling: ').bold = True
    p.add_run('Tailwind CSS 3\n')
    p.add_run('State Management: ').bold = True
    p.add_run('Zustand\n')
    p.add_run('Charts / Visualization: ').bold = True
    p.add_run('ECharts, Recharts, D3.js, deck.gl, Leaflet, Three.js\n')
    p.add_run('Other UI Libs: ').bold = True
    p.add_run('framer-motion, lucide-react, react-markdown, sonner, exceljs')

    doc.add_heading('AI & LLM Providers', level=2)
    p = doc.add_paragraph()
    p.add_run('Provider-switchable via environment variables:\n')
    p.add_run('- Default Local: Ollama (local GPU)\n')
    p.add_run('- Production / VM: Azure OpenAI\n')
    p.add_run('- Fast Inference: Groq\n')
    p.add_run('- Fallback: OpenRouter')

    doc.add_heading('Integrations', level=2)
    doc.add_paragraph('- Microsoft Graph (MSAL) for SharePoint file exchange\n- Oracle Primavera P6 Cloud (REST API + OAuth)\n- SAP via Excel extracts (ingested via pandas)\n- SAP BTP OData (Pulse Quality NCs/RFIs)\n- Transmission / Powerback APIs\n- Spectra Insights (Drone survey REST API)\n- SAP BTP UAT (E-Invoice)')

    doc.add_heading('3. Deployment Topology', level=1)
    doc.add_paragraph('The entire platform ships as a single process. FastAPI serves both the API and the built React Single Page Application (SPA).')
    p = doc.add_paragraph()
    p.add_run('- Web Traffic: ').bold = True
    p.add_run('Browser -> FastAPI (Uvicorn)\n')
    p.add_run('- Proxy: ').bold = True
    p.add_run('The application sits behind a reverse proxy at /akasha. API paths are internally rewritten.\n')
    p.add_run('- Corporate Proxy handling: ').bold = True
    p.add_run('Custom requests session patching to allow SSL traversal.\n')
    p.add_run('- Auto-migration: ').bold = True
    p.add_run('Runs dynamically on boot without Alembic migration files.')

    doc.save('Akasha_Tech_Stack.docx')

if __name__ == '__main__':
    generate_doc()
